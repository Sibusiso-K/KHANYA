"""REEFPRINT / KHANYA — 10-minute solution pitch, native PowerPoint.

Every number on a slide carries a provenance tag (LIVE APP, RECORDED, SOURCE, ASSUMPTION, ILLUSTRATIVE),
the deck-level version of the product's own rule: no number without its source.
Arithmetic is done here, in code, and printed, never typed by hand.
"""
import copy, os, sys, json
from PIL import Image, ImageDraw, ImageFilter
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn

D = os.path.dirname(os.path.abspath(__file__))
SCR = os.path.dirname(D)
IMG = os.path.join(D, "img")
CAP = os.path.join(SCR, "cap")
OUT = r"C:\Users\USER\Desktop\REEFPRINT\presentation\output"
VIDEO = r"C:\Users\USER\Desktop\REEFPRINT\presentation\video\REEFPRINT-promo-90s-embed.mp4"
MICRO = r"C:\Users\USER\Desktop\REEFPRINT\demo-images\test_11.jpg"
os.makedirs(OUT, exist_ok=True)
os.makedirs(os.path.join(D, "gen"), exist_ok=True)
GEN = os.path.join(D, "gen")

HEAD = "Georgia"
BODY = "Segoe UI"
BODYB = "Segoe UI Semibold"

NAVY = RGBColor(0x0E, 0x1B, 0x2E)
INK = RGBColor(0x1B, 0x2D, 0x40)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
LINE = RGBColor(0xD9, 0xDF, 0xE6)
PAPER = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COPPER = RGBColor(0xD2, 0x5A, 0x24)
PYRR = RGBColor(0x1F, 0xA3, 0xA8)
PENT = RGBColor(0x8E, 0x5C, 0xC9)
CHALC = RGBColor(0xE0, 0x7A, 0x2E)
MAG = RGBColor(0xC9, 0x9A, 0x1C)
GREEN = RGBColor(0x2E, 0x8B, 0x57)

TAGS = {  # provenance tag colours
    "LIVE APP": (PYRR, WHITE),
    "RECORDED": (INK, WHITE),
    "SOURCE": (RGBColor(0xE6, 0xEA, 0xEF), INK),
    "ASSUMPTION": (RGBColor(0xF6, 0xE3, 0xB4), RGBColor(0x6B, 0x4E, 0x00)),
    "ILLUSTRATIVE": (RGBColor(0x3A, 0x47, 0x57), WHITE),
    "SIMULATOR": (RGBColor(0xF6, 0xE3, 0xB4), RGBColor(0x6B, 0x4E, 0x00)),
    "ROADMAP": (RGBColor(0xE9, 0xDF, 0xF6), RGBColor(0x4B, 0x2A, 0x7A)),
}

# ---------------------------------------------------------------- numbers (computed, then printed)
N = {}
N["sa_pt_share"] = 120_000 / 170_000          # USGS MCS 2025, 2024e, kg
N["sa_pgm_reserve_share_max"] = 63_000_000 / 81_000_000  # world reserves ">81,000,000 kg"
BASKET = 45_993                                # Valterra H1 2026, R per PGM oz sold
TONNES = 250_000                               # ASSUMPTION: concentrator feed, t/month
GRADE = 4.0                                    # ASSUMPTION: head grade, g/t PGM
OZ = 31.1034768
N["uplift"] = {}
for pp in (0.1, 0.25, 0.5, 1.0):
    oz = TONNES * GRADE * (pp / 100) / OZ
    N["uplift"][pp] = (oz, oz * BASKET)
QEM = 1500                                     # SRC 2017 price list, per sample
from scipy.stats import beta
K_CONF, N_SEC = 4, 12                          # shipped refined pipeline: 1 continue + 3 grind finer answered; 8 hedged/refused
N["conf_rate"] = K_CONF / N_SEC
N["conf_lo"] = float(beta.ppf(0.025, K_CONF, N_SEC - K_CONF + 1))
N["conf_hi"] = float(beta.ppf(0.975, K_CONF + 1, N_SEC - K_CONF))
N["triage_point"] = 100 * N["conf_rate"] * QEM
N["triage_lo"] = 100 * N["conf_lo"] * QEM
N["triage_hi"] = 100 * N["conf_hi"] * QEM
LICENCE = 12_000
N["licence_vs_value"] = LICENCE / N["uplift"][0.5][1]
print(json.dumps({k: v for k, v in N.items()}, indent=1, default=str))

# ---------------------------------------------------------------- image helpers

def lens(src, out, size=900, crop=None, ring=None):
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
    w, h = im.size
    s = min(w, h)
    im = im.crop(((w - s) // 2, (h - s) // 2, (w - s) // 2 + s, (h - s) // 2 + s)).resize((size, size), Image.LANCZOS)
    m = Image.new("L", (size * 3, size * 3), 0)
    ImageDraw.Draw(m).ellipse([0, 0, size * 3 - 1, size * 3 - 1], fill=255)
    m = m.resize((size, size), Image.LANCZOS)
    im.putalpha(m)
    if ring:
        r = Image.new("RGBA", (size, size), (0, 0, 0, 0))
        ImageDraw.Draw(r).ellipse([6, 6, size - 7, size - 7], outline=ring + (255,), width=10)
        im = Image.alpha_composite(im, r)
    im.save(out)
    return out


def rounded(src, out, crop=None, radius=28, max_w=1600):
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    m = Image.new("L", (im.width * 2, im.height * 2), 0)
    ImageDraw.Draw(m).rounded_rectangle([0, 0, im.width * 2 - 1, im.height * 2 - 1], radius * 2, fill=255)
    m = m.resize(im.size, Image.LANCZOS)
    im.putalpha(m)
    im.save(out)
    return out


def darken(src, out, amount=0.55, tint=(14, 27, 46), crop=None):
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
    ov = Image.new("RGB", im.size, tint)
    Image.blend(im, ov, amount).save(out, quality=90)
    return out


G = {}
G["lens_title"] = lens(MICRO, os.path.join(GEN, "lens_title.png"), ring=(255, 255, 255))
G["lens_a"] = lens(MICRO, os.path.join(GEN, "lens_a.png"), crop=(800, 400, 2400, 2000))
G["lens_orig"] = lens(os.path.join(CAP, "hi_view_Original.png"), os.path.join(GEN, "lens_orig.png"), crop=(496, 0, 1860, 910))
G["lens_phase"] = lens(os.path.join(CAP, "hi_view_Phase_overlay.png"), os.path.join(GEN, "lens_phase.png"), crop=(496, 0, 1860, 910))
G["bg_title"] = darken(os.path.join(IMG, "aerial.jpg"), os.path.join(GEN, "bg_title.jpg"), 0.58)
G["bg_close"] = darken(os.path.join(IMG, "truckroad.jpg"), os.path.join(GEN, "bg_close.jpg"), 0.62)
G["haul"] = rounded(os.path.join(IMG, "haul.jpg"), os.path.join(GEN, "haul.png"))
G["plant"] = rounded(os.path.join(IMG, "plant.jpg"), os.path.join(GEN, "plant.png"), crop=(300, 0, 1920, 1080))
G["miner"] = rounded(os.path.join(IMG, "miner.jpg"), os.path.join(GEN, "miner.png"), crop=(200, 0, 1080, 720))
G["lab"] = rounded(os.path.join(IMG, "lab.jpg"), os.path.join(GEN, "lab.png"))
G["laptop"] = rounded(os.path.join(IMG, "laptop.png"), os.path.join(GEN, "laptop.png"), crop=(330, 0, 1590, 844))
G["phone"] = rounded(os.path.join(IMG, "phone.png"), os.path.join(GEN, "phone.png"), crop=(560, 0, 1360, 1080))
G["ui_overlay"] = rounded(os.path.join(CAP, "d03_ws_after.png"), os.path.join(GEN, "ui_overlay.png"), crop=(341, 323, 1823, 1053), radius=14)
G["ui_grain"] = rounded(os.path.join(CAP, "d05_grain4_top.png"), os.path.join(GEN, "ui_grain.png"), crop=(430, 600, 1823, 1080), radius=14)
G["ui_report"] = rounded(os.path.join(CAP, "d08_reports.png"), os.path.join(GEN, "ui_report.png"), crop=(96, 585, 1823, 770), radius=14)
G["ui_sim"] = rounded(os.path.join(CAP, "d11_process_sim.png"), os.path.join(GEN, "ui_sim.png"), crop=(655, 555, 1820, 735), radius=14)
G["ui_assist"] = rounded(os.path.join(CAP, "d18_assistant_answer.png"), os.path.join(GEN, "ui_assist.png"), crop=(1500, 110, 1900, 1060), radius=14)
G["ui_spatial"] = rounded(os.path.join(CAP, "d15_spatial_3D_model2.png"), os.path.join(GEN, "ui_spatial.png"), crop=(310, 380, 1560, 1080), radius=14)
G["ui_cand"] = rounded(os.path.join(CAP, "d09_candidate.png"), os.path.join(GEN, "ui_cand.png"), crop=(100, 0, 1820, 270), radius=14)
G["thumb_orig"] = rounded(os.path.join(CAP, "hi_view_Original.png"), os.path.join(GEN, "thumb_orig.png"), crop=(496, 0, 1860, 910), radius=18, max_w=700)
G["thumb_phase"] = rounded(os.path.join(CAP, "hi_view_Phase_overlay.png"), os.path.join(GEN, "thumb_phase.png"), crop=(496, 0, 1860, 910), radius=18, max_w=700)
G["thumb_grain"] = rounded(os.path.join(CAP, "d05_grain4_top.png"), os.path.join(GEN, "thumb_grain.png"), crop=(1545, 720, 1823, 1000), radius=10, max_w=700)
G["thumb_adv"] = rounded(os.path.join(CAP, "d10_process.png"), os.path.join(GEN, "thumb_adv.png"), crop=(655, 585, 1255, 692), radius=10, max_w=700)
G["thumb_sim"] = rounded(os.path.join(CAP, "d11_process_sim.png"), os.path.join(GEN, "thumb_sim.png"), crop=(1300, 585, 1810, 735), radius=10, max_w=700)

# ---------------------------------------------------------------- pptx helpers
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def morph(slide, dur=1100):
    """Morph transition (PowerPoint 2019+/365) with a fade fallback."""
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:p159="{P159}" Requires="p159">'
           f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           f'xmlns:p14="{P14}" spd="slow" p14:dur="{dur}"><p159:morph option="byObject"/></p:transition>'
           f'</mc:Choice><mc:Fallback><p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/></p:transition>'
           f'</mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    sld = slide._element
    clr = sld.find(qn("p:clrMapOvr"))
    if clr is not None:
        clr.addnext(el)
    else:
        sld.find(qn("p:cSld")).addnext(el)


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, name=None, line=None, italic=False):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        if line:
            p.line_spacing = line
        segs = para if isinstance(para, list) else [(para, {})]
        for seg, ov in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = ov.get("font", font)
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", italic)
            f.color.rgb = ov.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=PAPER, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if not shadow:
        sp = s._element.spPr
        ef = etree.SubElement(sp, qn("a:effectLst"))
    s.text_frame.text = ""
    return s


def tag(slide, x, y, label, name=None):
    fill, fg = TAGS[label]
    w = 0.24 + 0.092 * len(label)
    s = box(slide, x, y, w, 0.26, fill=fill, radius=0.13, name=name)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = BODYB
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = fg
    return s


def pic(slide, path, x, y, w=None, h=None, name=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if name:
        p.name = name
    return p


def eyebrow(slide, x, y, s, color=COPPER):
    text(slide, x, y, 8, 0.3, s.upper(), size=11, color=color, font=BODYB, bold=True)


def headline(slide, s, x=0.6, y=0.85, w=11.6, size=34, color=INK, name="!!headline"):
    return text(slide, x, y, w, 1.3, s, size=size, color=color, font=HEAD, bold=True, name=name, line=1.0)


def source(slide, s, y=7.05, color=MUTED):
    text(slide, 0.6, y, 12.1, 0.3, s, size=8.5, color=color)


def brand(slide, dark=False):
    c = WHITE if dark else INK
    text(slide, 0.6, 0.35, 4, 0.3, [[("REEFPRINT", {"font": BODYB, "bold": True, "size": 10.5, "color": c}),
                                    ("  ·  KHANYA", {"size": 10.5, "color": (RGBColor(0xB8, 0xC4, 0xD2) if dark else MUTED)})]],
         name="!!brand")


def pagenum(slide, n, dark=False):
    text(slide, 12.2, 0.35, 0.55, 0.3, f"{n:02d}", size=10.5, color=(RGBColor(0xB8, 0xC4, 0xD2) if dark else MUTED),
         align=PP_ALIGN.RIGHT, name="!!page")


def numcircle(slide, x, y, n, color=INK, d=0.42):
    c = box(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name = BODYB
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = WHITE
    return c


def stat(slide, x, y, big, label, tg, w=3.6, color=INK, big_size=40):
    text(slide, x, y, w, 0.75, big, size=big_size, color=color, font=HEAD, bold=True)
    text(slide, x, y + 0.78, w, 0.75, label, size=12.5, color=INK, line=1.05)
    tag(slide, x, y + 1.55 if len(label) > 70 else y + 1.4, tg)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def style_chart(chart, colors, legend=False, size=11):
    chart.font.name = BODY
    chart.font.size = Pt(size)
    chart.font.color.rgb = INK
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
    for s, c in zip(chart.plots[0].series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c


# ================================================================ SLIDES
S = []


def new(dark=False):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY if dark else WHITE)
    S.append(s)
    return s

# ---- 1 TITLE
s = new(dark=True)
pic(s, G["bg_title"], 0, 0, w=SW, h=SH)
pic(s, G["lens_title"], 8.2, 1.15, w=4.4, name="!!lens")
eyebrow(s, 0.7, 1.35, "Mintek–SCi Grad Hackathon 2026  ·  Problem 3", color=RGBColor(0xF2, 0x9A, 0x6A))
text(s, 0.7, 1.8, 7.6, 1.3, "REEFPRINT", size=72, color=WHITE, font=HEAD, bold=True, name="!!headline")
text(s, 0.72, 3.05, 7.4, 0.5, "otherwise known as KHANYA", size=22, color=RGBColor(0xC9, 0xD5, 0xE3), font=HEAD, italic=True)
text(s, 0.72, 3.85, 7.0, 1.0, "Real-time mineral intelligence from a micrograph — and the honesty to say when it doesn't know.",
     size=19, color=WHITE, line=1.1)
text(s, 0.72, 5.45, 7.4, 0.6, [[("Team Sonar", {"font": BODYB, "bold": True, "color": WHITE}),
                                ("   Lethabo Hoaeane  ·  Sibusiso Khumalo  ·  Ipeleng Modise", {"color": RGBColor(0xC9, 0xD5, 0xE3)})]], size=14)
text(s, 0.72, 5.85, 7.4, 0.4, "Unisa  ·  Wits  ·  TUT", size=12, color=RGBColor(0x9F, 0xAE, 0xBF))
tag(s, 0.72, 6.85, "ILLUSTRATIVE")
text(s, 2.1, 6.86, 6, 0.3, "Background: Mixkit stock footage. Lens: LumenStone S2 test_11, public reflected-light micrograph.", size=8.5,
     color=RGBColor(0x9F, 0xAE, 0xBF))
morph(s)
notes(s, "SCRIPT_1")

# ---- 2 PROBLEM
s = new()
brand(s); pagenum(s, 2)
eyebrow(s, 0.6, 0.85, "The problem")
headline(s, "The plant decides every hour.\nThe mineralogy arrives days later.", y=1.15, size=36)
stat(s, 0.6, 2.85, "Days", "Laboratory SEM or XRD phase characterisation “can take days” — the challenge brief itself.", "SOURCE", w=3.4)
stat(s, 4.15, 2.85, "1 week", "is the expedited turnaround for automated mineralogy at a major commercial lab; standard is longer.", "SOURCE", w=3.4)
stat(s, 7.7, 2.85, "$1,500", "per QEMSCAN modal-mineralogy and liberation analysis on one published lab price list.", "SOURCE", w=3.0)
pic(s, G["lens_a"], 10.75, 2.55, w=2.0, name="!!lens")
box(s, 0.6, 5.2, 12.1, 1.45, fill=PAPER, radius=0.12)
text(s, 0.95, 5.38, 11.4, 1.2, [[("Consequence, in the brief's words: ", {"font": BODYB, "bold": True}),
                                 ("without real-time feedback, processing plants cannot adjust to changes in ore quality, leading to inefficient chemical usage and lower mineral yields. ",
                                  {}),
                                 ("Ore changes by the hour; the answer about it comes back after that ore has already been milled, floated and lost.", {"italic": True})]],
     size=15, line=1.15)
source(s, "Sources: Mintek–SCi Grad Hackathon 2026, Problem 3 brief · ALS Global, Mineralogy turnaround FAQ · Saskatchewan Research Council, Advanced Microanalysis Centre, QEMSCAN® price list (2017).")
morph(s)
notes(s, "SCRIPT_2")

# ---- 3 WHO IT AFFECTS
s = new()
brand(s); pagenum(s, 3)
eyebrow(s, 0.6, 0.85, "Who it affects")
headline(s, "South Africa's economy runs through\nthese flotation cells.", y=1.15, size=34)
stat(s, 0.6, 2.75, f"{N['sa_pt_share']*100:.0f}%", "of the world's mined platinum came from South Africa in 2024 (120 t of 170 t).", "SOURCE", w=3.2, color=PYRR)
stat(s, 0.6, 4.55, "≈ ¾", "of the world's PGM reserves are South African — 63,000 t of >81,000 t, mostly in the Bushveld Complex.", "SOURCE", w=3.2, color=PENT)
stat(s, 4.2, 2.75, "6.1%", "of South Africa's nominal GDP came from mining in 2024.", "SOURCE", w=3.0, color=CHALC)
stat(s, 4.2, 4.55, "474,736", "people directly employed in mining in 2024.", "SOURCE", w=3.0, color=INK)
# who, as a column of rows
rows = [("Metallurgists", "set grind, reagent and air without knowing today's ore"),
        ("Mineralogists & labs", "a queue of samples, most of them routine"),
        ("Operators", "react to recovery after it has already dropped"),
        ("Producers", "lose payable metal to tailings, overspend reagent and power"),
        ("Communities & fiscus", "jobs, royalties and tax depend on margin per ounce")]
box(s, 7.75, 2.65, 5.0, 3.95, fill=PAPER, radius=0.12)
text(s, 8.05, 2.82, 4.5, 0.35, "WHO FEELS IT", size=10.5, color=COPPER, font=BODYB, bold=True)
for i, (a, b) in enumerate(rows):
    yy = 3.22 + i * 0.66
    box(s, 8.05, yy + 0.08, 0.16, 0.16, fill=[PYRR, PENT, CHALC, MAG, INK][i], shape=MSO_SHAPE.OVAL)
    text(s, 8.35, yy, 4.2, 0.62, [[(a + "  ", {"font": BODYB, "bold": True}), (b, {"color": MUTED})]], size=12, line=1.05)
pic(s, G["lens_phase"], 11.65, 0.75, w=1.2, name="!!lens")
source(s, "Sources: USGS Mineral Commodity Summaries 2025, Platinum-group metals (2024 estimates; world reserves stated as >81,000,000 kg) · Department of Mineral and Petroleum Resources, Mining Sector Performance 2024.")
morph(s)
notes(s, "SCRIPT_3")

# ---- 4 HOW IT HURTS
s = new()
brand(s); pagenum(s, 4)
eyebrow(s, 0.6, 0.85, "How it hurts")
headline(s, "When ore changes faster than the lab,\nevery plant setting is a guess.", y=1.15, size=34)
steps = [("Ore changes", "harder, finer, more\nlocked sulphides", PYRR),
         ("No feedback", "mineralogy is days\naway", MUTED),
         ("Grind guessed", "too fine wastes power;\ntoo coarse locks metal", CHALC),
         ("Reagents guessed", "dosed for yesterday's\nore", MAG),
         ("Metal lost", "locked grains report\nto tailings", PENT)]
for i, (a, b, c) in enumerate(steps):
    x = 0.6 + i * 2.5
    box(s, x, 2.85, 2.2, 1.85, fill=PAPER, radius=0.14)
    box(s, x + 0.25, 3.05, 0.34, 0.34, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.25, 3.5, 1.85, 0.4, a, size=15, font=BODYB, bold=True)
    text(s, x + 0.25, 3.9, 1.9, 0.75, b, size=11.5, color=MUTED, line=1.05)
    if i < 4:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.24), Inches(3.62), Inches(0.24), Inches(0.3))
        ar.fill.solid(); ar.fill.fore_color.rgb = LINE; ar.line.fill.background()
pic(s, G["plant"], 0.6, 5.0, h=1.85)
text(s, 4.2, 5.15, 8.4, 1.6, [[("The lever exists — grind, reagent, air. ", {"font": BODYB, "bold": True}),
                              ("What is missing is a fast, trustworthy reading of the minerals that tells the plant which way to move it, and when ", {}),
                              ("not", {"italic": True}), (" to.", {})]], size=16, line=1.15)
tag(s, 0.6, 6.95, "ILLUSTRATIVE")
pic(s, G["lens_a"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_4")

# ---- 5 SOLUTION
s = new()
brand(s); pagenum(s, 5)
eyebrow(s, 0.6, 0.85, "The solution")
headline(s, "REEFPRINT turns a micrograph into a\nplant decision — on a laptop, in minutes.", y=1.15, size=32)
pic(s, G["laptop"], 7.05, 2.55, w=5.7)
tag(s, 7.05, 6.4, "LIVE APP")
text(s, 8.3, 6.41, 4.4, 0.3, "Real REEFPRINT dashboard, keyed into stock laptop footage.", size=8.5, color=MUTED)
cards = [("1", "Identify mineral phases", "Pyrrhotite, pentlandite and chalcopyrite mapped pixel by pixel, with area fractions. Magnetite attempted and reported.", PYRR, "LIVE APP"),
         ("2", "Accuracy report", "Bound to the exact checkpoint hash, on 12 held-out sections, zeros included.", INK, "RECORDED"),
         ("3", "Adjust a plant parameter", "Advisory → simulated setpoint over OPC UA. Held automatically when the model isn't approved.", CHALC, "SIMULATOR")]
for i, (n, a, b, c, tg) in enumerate(cards):
    y = 2.55 + i * 1.32
    box(s, 0.6, y, 6.1, 1.18, fill=PAPER, radius=0.12)
    numcircle(s, 0.82, y + 0.22, n, color=c)
    text(s, 1.45, y + 0.14, 3.9, 0.4, a, size=16, font=BODYB, bold=True)
    tag(s, 5.0, y + 0.18, tg)
    text(s, 1.45, y + 0.52, 5.0, 0.62, b, size=11.5, color=MUTED, line=1.05)
text(s, 0.6, 6.6, 6.2, 0.4, [[("Plus: ", {"font": BODYB, "bold": True}),
                             ("grain liberation (free vs locked) · XRF/assay context · spatial view · evidence companion · phone access", {"color": MUTED})]],
     size=11)
pic(s, G["lens_phase"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_5")

# ---- 6 HOW IT WORKS
s = new()
brand(s); pagenum(s, 6)
eyebrow(s, 0.6, 0.85, "How it works")
headline(s, "One learned stage. Everything after it is\ndeterministic, auditable — and allowed to refuse.", y=1.15, size=30)
flow = [("Polished section", "standard ore-microscopy\nprep", MUTED),
        ("Reflected-light image", "any lab microscope\ncamera", MUTED),
        ("Quality gate", "refuses blur, glare,\nwrong lighting", COPPER),
        ("Segment phases", "DeepLabV3 / ResNet-50,\n6 fields, laptop CPU", PYRR),
        ("Grains & liberation", "watershed particles;\nfree vs locked", PENT),
        ("Uncertainty", "conformal band around\neach decision floor", MAG),
        ("Advise or refuse", "act · verify · hold,\nwith a stated reason", CHALC),
        ("Setpoint + record", "OPC UA, only if\napproved; full audit", INK)]
for i, (a, b, c) in enumerate(flow):
    x = 0.6 + i * 1.53
    box(s, x, 2.6, 1.38, 1.75, fill=PAPER, radius=0.12)
    box(s, x + 0.14, 2.75, 0.3, 0.3, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.14, 3.13, 1.15, 0.5, a, size=11.5, font=BODYB, bold=True, line=0.95)
    text(s, x + 0.14, 3.62, 1.2, 0.7, b, size=9, color=MUTED, line=1.0)
pic(s, G["thumb_orig"], 0.6, 4.6, w=2.6)
text(s, 0.6, 6.37, 2.6, 0.3, "Input · test_11", size=9.5, color=MUTED)
pic(s, G["thumb_phase"], 3.4, 4.6, w=2.6)
text(s, 3.4, 6.37, 2.6, 0.3, "Phase map", size=9.5, color=MUTED)
pic(s, G["thumb_grain"], 6.2, 4.6, h=1.72)
text(s, 6.2, 6.37, 1.9, 0.3, "Grain 4 · FREE", size=9.5, color=MUTED)
pic(s, G["thumb_adv"], 8.15, 4.6, w=4.6)
text(s, 8.15, 5.45, 4.6, 0.3, "Advisory: “Marginal — verify before acting”", size=9.5, color=MUTED)
pic(s, G["thumb_sim"], 8.15, 5.78, w=2.9)
text(s, 11.15, 5.95, 1.6, 0.6, "Simulator: setting held", size=9.5, color=MUTED)
tag(s, 0.6, 6.95, "LIVE APP")
text(s, 1.72, 6.96, 10, 0.3, "Thumbnails are real outputs from the running workbench, 1 Oct 2026, active checkpoint fb78727d.", size=8.5, color=MUTED)
pic(s, G["lens_orig"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_6")

# ---- 7 DEMO VIDEO
s = new(dark=True)
brand(s, dark=True); pagenum(s, 7, dark=True)
eyebrow(s, 0.6, 0.62, "Demo  ·  90 seconds", color=RGBColor(0xF2, 0x9A, 0x6A))
text(s, 0.6, 0.95, 11, 0.6, "From the mine to a decision, on the real app.", size=26, color=WHITE, font=HEAD, bold=True, name="!!headline")
poster = os.path.join(GEN, "poster.jpg")
if os.path.exists(VIDEO):
    import subprocess
    FF = r"C:\Users\USER\Desktop\REEFPRINT\.workbench\media-tools\imageio_ffmpeg\binaries\ffmpeg-win-x86_64-v7.1.exe"
    subprocess.run([FF, "-v", "error", "-y", "-ss", "21", "-i", VIDEO, "-frames:v", "1", "-q:v", "2", poster], check=True)
    mv = s.shapes.add_movie(VIDEO, Inches(1.72), Inches(1.68), Inches(9.9), Inches(5.569), poster_frame_image=poster, mime_type="video/mp4")
    mv.name = "demo_video"
else:
    box(s, 1.27, 1.75, 10.8, 5.0, fill=INK)
    text(s, 1.27, 4.0, 10.8, 0.5, "[demo video not yet rendered]", size=18, color=WHITE, align=PP_ALIGN.CENTER)
morph(s)
notes(s, "SCRIPT_7")

# ---- 8 EVIDENCE
s = new()
brand(s); pagenum(s, 8)
eyebrow(s, 0.6, 0.85, "Evidence, including what fails")
headline(s, "We report what the model gets wrong — first.", y=1.15, size=32)
cd = CategoryChartData()
cd.categories = ["Chalcopyrite", "Pyrrhotite", "Pentlandite", "Magnetite"]
cd.add_series("Active model fb78727d (live)", (0.354, 0.687, 0.355, 0.000))
cd.add_series("Retrained candidate 42646cfa (quarantined)", (0.7998, 0.8183, 0.4460, 0.2477))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2.0), Inches(7.2), Inches(4.75), cd)
ch = gf.chart
style_chart(ch, [INK, PYRR], legend=True, size=11)
ch.has_title = True
ch.chart_title.text_frame.text = "Intersection-over-Union per phase, 12 held-out LumenStone S2 sections"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
pl = ch.plots[0]
pl.gap_width = 70
pl.has_data_labels = True
pl.data_labels.number_format = '0.00'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(10)
va = ch.value_axis
va.maximum_scale = 1.0
va.minimum_scale = 0
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = LINE
va.tick_labels.font.size = Pt(10)
va.format.line.fill.background()
ch.category_axis.tick_labels.font.size = Pt(11)
ch.category_axis.format.line.color.rgb = LINE
x0 = 8.05
text(s, x0, 2.1, 4.6, 0.6, "0.454", size=40, font=HEAD, bold=True)
text(s, x0 + 1.75, 2.28, 2.9, 0.6, "mean IoU, live model · pixel accuracy 77.2% · magnetite 0", size=11.5, color=MUTED, line=1.05)
tag(s, x0, 2.95, "RECORDED")
text(s, x0, 3.35, 4.6, 0.6, "0.632", size=40, font=HEAD, bold=True, color=PYRR)
text(s, x0 + 1.75, 3.53, 2.9, 0.6, "candidate · finds 89% of magnetite but only 25.5% precise → not deployed", size=11.5, color=MUTED, line=1.05)
tag(s, x0, 4.2, "RECORDED")
text(s, x0, 4.6, 4.6, 0.6, "98 s", size=40, font=HEAD, bold=True, color=CHALC)
text(s, x0 + 1.75, 4.78, 2.9, 0.6, "fresh six-field analysis on a laptop CPU, no GPU, no cloud", size=11.5, color=MUTED, line=1.05)
tag(s, x0, 5.45, "LIVE APP")
box(s, x0, 5.85, 4.65, 0.95, fill=PAPER, radius=0.1)
text(s, x0 + 0.2, 5.93, 4.3, 0.85, [[("Refusal by design: ", {"font": BODYB, "bold": True}),
                                     ("0 unsafe advisories against expert-mask advice on 12 held-out sections (raw pipeline: 2). Every disagreement became “verify”. 0 of 12 is an observation (95% upper bound 22%), not a guarantee.", {"color": MUTED})]],
     size=10, line=1.05)
source(s, "Scope: LumenStone S2 (Norilsk Ni-Cu-PGE sulphide) — an assemblage analogue for Bushveld base-metal sulphides, not UG2; no chromite data yet. Candidate scores are a fixed regression check on historically reused test images. Refusal result: historical checkpoint, reports/ACCURACY-REPORT.md §7.")
pic(s, G["lens_phase"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_8")

# ---- 9 COMPETITION
s = new()
brand(s); pagenum(s, 9)
eyebrow(s, 0.6, 0.85, "Competition and what sets us apart")
headline(s, "Everyone measures minerals. We also measure\nwhether the answer is safe to act on.", y=1.15, size=30)
cols = ["", "Mineral phases", "Liberation", "Time to answer", "Where it runs", "Refuses with a reason", "Drives a setpoint"]
rowsd = [("SEM automated mineralogy\nQEMSCAN · MLA · TIMA · Mineralogic", "✓", "✓", "days – weeks", "central lab", "—", "—"),
         ("XRD", "✓ bulk", "—", "hours – days", "lab", "—", "—"),
         ("Handheld / online XRF", "elements only", "—", "seconds", "field, belt", "—", "via models"),
         ("Hyperspectral ore sensing", "partial*", "—", "real time", "belt, core", "—", "partial"),
         ("Flotation control, e.g. Mintek FloatStar", "—", "—", "real time", "plant", "—", "✓"),
         ("REEFPRINT / KHANYA", "✓ 3 phases", "✓ 2D", "≈ 1.5 min", "laptop, lab or plant", "✓", "✓ gated")]
tbl = s.shapes.add_table(len(rowsd) + 1, len(cols), Inches(0.6), Inches(2.45), Inches(12.1), Inches(3.9)).table
widths = [3.3, 1.45, 1.2, 1.55, 1.75, 1.55, 1.3]
for i, wv in enumerate(widths):
    tbl.columns[i].width = Inches(wv)
for r in range(len(rowsd) + 1):
    for c in range(len(cols)):
        cell = tbl.cell(r, c)
        val = cols[c] if r == 0 else rowsd[r - 1][c]
        cell.text = ""
        tf = cell.text_frame
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = val.split("\n")
        for li, ln in enumerate(lines):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            rr = p.add_run()
            rr.text = ln
            ours = r == len(rowsd)
            rr.font.name = BODYB if (r == 0 or ours or (c == 0 and li == 0)) else BODY
            rr.font.bold = r == 0 or ours or (c == 0 and li == 0)
            rr.font.size = Pt(10.5 if li == 0 else 8.5)
            rr.font.color.rgb = WHITE if (r == 0 or ours) else (INK if li == 0 else MUTED)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (PYRR if r == len(rowsd) else (WHITE if r % 2 else PAPER))
text(s, 0.6, 6.45, 12.1, 0.55, [[("Complementary, not a replacement: ", {"font": BODYB, "bold": True}),
                                ("REEFPRINT routes only the uncertain samples to QEMSCAN, and hands checked setpoints to the controllers plants already run.", {})]],
     size=13)
source(s, "Capability summary from public product descriptions; “—” means not that tool's role. *SWIR hyperspectral identifies minerals by vibrational absorption features, which opaque phases such as chromite lack.", y=7.0)
pic(s, G["lens_a"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_9")

# ---- 10 VALUE
s = new()
brand(s); pagenum(s, 10)
eyebrow(s, 0.6, 0.85, "Why a PGM producer wants this today")
headline(s, "Spend the expensive lab on the hard samples.\nMove the plant inside the shift.", y=1.15, size=30)
cd = CategoryChartData()
cd.categories = ["+0.1 pp", "+0.25 pp", "+0.5 pp", "+1.0 pp"]
cd.add_series("Rand per month", tuple(round(N["uplift"][k][1] / 1e6, 2) for k in (0.1, 0.25, 0.5, 1.0)))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.75), Inches(2.45), Inches(6.0), Inches(3.75), cd)
ch = gf.chart
style_chart(ch, [PENT], legend=False)
ch.has_title = True
ch.chart_title.text_frame.text = "Value of a recovery change, R million / month"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.number_format = 'R0.0"m"'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(11)
pl.data_labels.font.bold = True
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
ch.category_axis.tick_labels.font.size = Pt(11)
ch.category_axis.format.line.color.rgb = LINE
tag(s, 6.75, 6.3, "ILLUSTRATIVE")
tag(s, 8.2, 6.3, "ASSUMPTION")
text(s, 6.75, 6.62, 6.0, 0.45, f"If a site-validated advisory moved recovery by these amounts: {TONNES:,} t/month at {GRADE:g} g/t (assumed), R{BASKET:,}/PGM oz (Valterra H1 2026 basket). We have not measured any recovery change.",
     size=8.5, color=MUTED, line=1.0)
levers = [("Triage the lab queue", f"A third of held-out sections got a confident optical answer (4/12; 95% CI {N['conf_lo']*100:.0f}–{N['conf_hi']*100:.0f}%); the rest were sent for review. Per 100 samples that avoids ≈ ${N['triage_point']:,.0f} of QEMSCAN time (range ${N['triage_lo']:,.0f}–${N['triage_hi']:,.0f}).", PYRR, "ASSUMPTION"),
          ("Decide inside the shift", "≈ 1.5 minutes per section on a laptop, against days in the lab queue — before the ore it describes has been floated.", CHALC, "LIVE APP"),
          ("Never act on thin evidence", "A conservative default with a reason instead of a guess: the costliest error at an ore transition is a confident wrong setpoint.", INK, "RECORDED")]
for i, (a, b, c, tg) in enumerate(levers):
    y = 2.45 + i * 1.42
    box(s, 0.6, y, 5.85, 1.28, fill=PAPER, radius=0.12)
    numcircle(s, 0.8, y + 0.2, i + 1, color=c)
    text(s, 1.4, y + 0.13, 3.7, 0.35, a, size=15, font=BODYB, bold=True)
    tag(s, 5.0, y + 0.16, tg)
    text(s, 1.4, y + 0.5, 4.9, 0.78, b, size=10.5, color=MUTED, line=1.03)
source(s, "QEMSCAN price: SRC Advanced Microanalysis Centre list (2017). Confident-answer rate: shipped decision-gap run, reports/decision_gap_patches_refined.json (exact Clopper–Pearson interval). Basket price: Valterra Platinum H1 2026 short-form results, rand basket price per PGM ounce sold.")
pic(s, G["lens_phase"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_10")

# ---- 11 BUSINESS MODEL, ACCESS, MARKET
s = new()
brand(s); pagenum(s, 11)
eyebrow(s, 0.6, 0.85, "Business model, access and market")
headline(s, "Priced like plant software.\nProven like a laboratory method.", y=1.15, size=32)
tiers = [("Paid pilot", "R210,000", "8–12 weeks · one lab or concentrator · shadow mode · go/no-go gates", PYRR),
         ("Site licence", "R12,000/mo", "per site, annual contract · + R30k–R75k one-off deployment and integration", CHALC),
         ("Enterprise", "Quoted", "multi-site · on-prem model registry · custom ore domains · SLA and support", PENT)]
for i, (a, price, b, c) in enumerate(tiers):
    x = 0.6 + i * 2.75
    box(s, x, 2.5, 2.55, 2.45, fill=PAPER, radius=0.14)
    box(s, x + 0.22, 2.7, 0.3, 0.3, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.22, 3.1, 2.2, 0.35, a, size=14, font=BODYB, bold=True)
    text(s, x + 0.22, 3.45, 2.3, 0.5, price, size=20, font=HEAD, bold=True, color=c)
    text(s, x + 0.22, 4.0, 2.15, 0.9, b, size=10, color=MUTED, line=1.03)
tag(s, 0.6, 5.08, "ASSUMPTION")
text(s, 1.95, 5.09, 6.3, 0.5, f"Price hypotheses to test with buyers, not observed prices. At the illustrative +0.5 pp, a site licence is ≈ {N['licence_vs_value']*100:.2f}% of the value it would protect.",
     size=9.5, color=MUTED, line=1.03)
text(s, 0.6, 5.6, 7.9, 0.35, "HOW CUSTOMERS GET IT", size=10.5, color=COPPER, font=BODYB, bold=True)
text(s, 0.6, 5.92, 7.9, 1.0, [[("Offline workstation ", {"font": BODYB, "bold": True}), ("— runs on a lab laptop with no internet.  ", {"color": MUTED}),
                              ("Private web workspace ", {"font": BODYB, "bold": True}), ("— sign-in, phone and tablet (PWA).  ", {"color": MUTED}),
                              ("Plant connector ", {"font": BODYB, "bold": True}), ("— OPC UA to SCADA / flotation control, behind approval gates.", {"color": MUTED})]],
     size=11, line=1.1)
box(s, 8.85, 2.5, 3.9, 4.4, fill=NAVY, radius=0.14)
text(s, 9.1, 2.7, 3.5, 0.35, "WHO BUYS FIRST", size=10.5, color=RGBColor(0xF2, 0x9A, 0x6A), font=BODYB, bold=True)
text(s, 9.1, 3.05, 3.5, 2.2, ["South African PGM and chrome producers:",
                             [("Valterra Platinum · Impala Platinum · Sibanye-Stillwater · Northam Platinum · African Rainbow Minerals · Tharisa", {"font": BODYB, "bold": True})],
                             "Mineralogy labs as triage partners:",
                             [("Mintek · commercial labs", {"font": BODYB, "bold": True})],
                             "Then base metals, chrome and iron ore across Africa."], size=11, color=WHITE, spacing=6, line=1.05)
text(s, 9.1, 5.7, 3.5, 0.5, "Daily users: metallurgists, mineralogists, control-room operators; geologists view results.", size=10, color=RGBColor(0xC9, 0xD5, 0xE3), line=1.05)
text(s, 9.1, 6.35, 3.5, 0.4, "Named as target segments — not customers or endorsements.", size=8.5, color=RGBColor(0x9F, 0xAE, 0xBF))
pic(s, G["lens_a"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_11")

# ---- 12 ARCHITECTURE & SCALE
s = new()
brand(s); pagenum(s, 12)
eyebrow(s, 0.6, 0.85, "Technology and scalability")
headline(s, "Boring, permissive, offline-first technology.\nScales by adding sites, not servers.", y=1.15, size=30)
layers = [("People", "browser · phone (PWA) · control room", MUTED, 0.6),
          ("React 19 + TypeScript + Vite", "workbench UI · three.js 3D · offline-capable", PYRR, 2.45),
          ("FastAPI service (Python)", "jobs · evidence · reports · approval gates · assistant", INK, 4.3),
          ("Model worker", "PyTorch DeepLabV3/ResNet-50 · OpenCV grains · conformal bands", PENT, 6.15),
          ("Plant connector", "OPC UA (asyncua) · simulator first · SCADA / FloatStar later", CHALC, 8.0)]
for a, b, c, x in layers:
    box(s, x, 2.6, 1.7, 2.0, fill=PAPER, radius=0.12)
    box(s, x + 0.17, 2.77, 0.3, 0.3, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.17, 3.15, 1.4, 0.7, a, size=12, font=BODYB, bold=True, line=0.95)
    text(s, x + 0.17, 3.85, 1.42, 0.75, b, size=9, color=MUTED, line=1.0)
for x in (2.32, 4.17, 6.02, 7.87):
    ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(3.45), Inches(0.12), Inches(0.24))
    ar.fill.solid(); ar.fill.fore_color.rgb = LINE; ar.line.fill.background()
box(s, 0.6, 4.8, 9.1, 0.75, fill=PAPER, radius=0.1)
text(s, 0.8, 4.9, 8.8, 0.6, [[("Data: ", {"font": BODYB, "bold": True}), ("Supabase Postgres + private object storage, tenant-isolated; every result carries model, checkpoint SHA-256, source and runtime. ", {"color": MUTED}),
                             ("Remote access: ", {"font": BODYB, "bold": True}), ("Cloudflare tunnel, sign-in required.", {"color": MUTED})]], size=10.5, line=1.05)
text(s, 0.6, 5.75, 9.1, 1.2, [[("Scales because: ", {"font": BODYB, "bold": True}),
                              ("inference runs on CPU at the site (GPU optional), each site is an isolated worker with its own ore model, the cloud is optional, and the whole stack is permissively licensed with an SBOM — assignable to Mintek. ", {"color": MUTED}),
                              ("Quality is guarded by 72 backend checks and Playwright browser tests in CI.", {"color": MUTED})]], size=11, line=1.1)
pic(s, G["phone"], 10.05, 2.45, h=4.45)
tag(s, 10.05, 6.98, "LIVE APP")
pic(s, G["lens_orig"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_12")

# ---- 13 MULTI-MODEL ENGINE
s = new()
brand(s); pagenum(s, 13)
eyebrow(s, 0.6, 0.85, "How accuracy keeps climbing")
headline(s, "Different models are good at different minerals.\nSo we route each answer to the best expert.", y=1.15, size=29)
# diagram
nodes = [("Image in", 0.6, 3.55, MUTED), ("Quality gate\n+ ore-domain check", 2.25, 3.55, COPPER), ("Router", 4.15, 3.55, INK)]
for a, x, y, c in nodes:
    box(s, x, y, 1.45, 1.0, fill=PAPER, radius=0.12)
    box(s, x + 0.14, y + 0.14, 0.24, 0.24, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.14, y + 0.43, 1.25, 0.55, a, size=10.5, font=BODYB, bold=True, line=0.95)
experts = [("General 5-phase model", "pyrrhotite · chalcopyrite", PYRR), ("Magnetite specialist", "recall 89% today, precision to fix", MAG),
           ("Pentlandite / BMS expert", "best historical 0.547", PENT), ("UG2 + chromite model", "needs Bushveld labels", CHALC),
           ("Polarimetry model", "Stokes channels (roadmap)", INK)]
for i, (a, b, c) in enumerate(experts):
    y = 2.35 + i * 0.68
    box(s, 5.95, y, 2.6, 0.58, fill=PAPER, radius=0.1)
    box(s, 6.07, y + 0.17, 0.22, 0.22, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, 6.38, y + 0.04, 2.15, 0.3, a, size=10, font=BODYB, bold=True)
    text(s, 6.38, y + 0.3, 2.15, 0.26, b, size=8.5, color=MUTED)
box(s, 8.85, 3.2, 1.75, 1.7, fill=PAPER, radius=0.12)
text(s, 9.0, 3.32, 1.5, 1.5, [[("Per-phase weighted vote", {"font": BODYB, "bold": True})], [("weights fitted on validation data only — never the test set", {"color": MUTED, "size": 8.5})]], size=10.5, line=1.0, spacing=3)
box(s, 10.85, 3.2, 1.9, 1.7, fill=NAVY, radius=0.12)
text(s, 11.0, 3.32, 1.65, 1.5, [[("Referee", {"font": BODYB, "bold": True, "color": WHITE})],
                                [("conformal check: confident → answer; disagree → “verify”; thin evidence → hold", {"color": RGBColor(0xC9, 0xD5, 0xE3), "size": 8.5})]], size=11, line=1.0, spacing=3)
for x in (2.07, 3.97, 5.62, 8.62, 10.65):
    ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(3.95), Inches(0.16), Inches(0.22))
    ar.fill.solid(); ar.fill.fore_color.rgb = LINE; ar.line.fill.background()
tag(s, 0.6, 5.95, "RECORDED")
text(s, 1.95, 5.92, 10.8, 0.5, [[("Our own data already shows it: ", {"font": BODYB, "bold": True}),
                                ("the live model never finds magnetite; the candidate finds 89% of it; the historical model is best on pyrrhotite (0.870) and pentlandite (0.547).", {"color": MUTED})]],
     size=11, line=1.05)
tag(s, 0.6, 6.5, "ROADMAP")
text(s, 1.85, 6.47, 10.9, 0.5, [[("Continuous training loop: ", {"font": BODYB, "bold": True}),
                                ("every QEMSCAN-labelled section retrains the specialists → shadow test → promote only if validation and conformal coverage both hold → model registry by SHA.", {"color": MUTED})]],
     size=11, line=1.05)
pic(s, G["lens_phase"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_13")

# ---- 14 ROADMAP & HYPOTHESES
s = new()
brand(s); pagenum(s, 14)
eyebrow(s, 0.6, 0.85, "Pathway forward")
headline(s, "From a working demo to every concentrator.", y=1.15, size=32)
phases = [("Today", "Oct 2026", "Working app · 3 phases · accuracy report · gated simulator · offline", PYRR),
          ("0–3 months", "Lab shadow pilot", "R210k · QEMSCAN-labelled sections from a partner lab · measured turnaround", CHALC),
          ("3–9 months", "Advisory pilot", "Specialist ensemble · UG2 + chromite domain · one concentrator, metallurgist approves", PENT),
          ("9–18 months", "Closed loop", "Controlled flotation trial with plant control · polarimetry rig · side-stream imaging", MAG),
          ("18 months +", "Product", "Multi-site service across PGM, chrome, base metals and iron ore", INK)]
ln = s.shapes.add_connector(1, Inches(0.8), Inches(2.62), Inches(12.6), Inches(2.62))
ln.line.color.rgb = LINE
ln.line.width = Pt(2.5)
for i, (a, b, c, col) in enumerate(phases):
    x = 0.6 + i * 2.45
    box(s, x + 0.05, 2.47, 0.3, 0.3, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, x, 2.9, 2.3, 0.3, a, size=11, color=col, font=BODYB, bold=True)
    text(s, x, 3.18, 2.3, 0.4, b, size=15, font=HEAD, bold=True)
    text(s, x, 3.6, 2.2, 0.9, c, size=10, color=MUTED, line=1.03)
box(s, 0.6, 4.65, 12.15, 2.25, fill=PAPER, radius=0.12)
text(s, 0.85, 4.8, 6, 0.3, "HYPOTHESES WE WILL TEST, AND HOW THEY COULD FAIL", size=10.5, color=COPPER, font=BODYB, bold=True)
hyps = [("H1", "QEMSCAN-labelled training lifts magnetite from 0 towards the 0.65 the dataset's authors report.", "Fails if the reflectance contrast is too low, even with better labels."),
        ("H2", "Polarisation channels split pentlandite (isotropic) from pyrrhotite (anisotropic) better than colour alone.", "Maths proven on a synthetic phantom (40× separation); the real-data test is still outstanding."),
        ("H3", "The refusal gate keeps unsafe advisories at zero while the answered share rises above 50%.", "Fails if coverage stays low on site ore — then we sell triage, not control."),
        ("H4", "Optical triage lowers mineralogy cost per plant decision.", "Measured in the pilot as cost and turnaround per approved decision.")]
for i, (h, a, b) in enumerate(hyps):
    x = 0.85 + (i % 2) * 6.0
    y = 5.15 + (i // 2) * 0.85
    text(s, x, y, 0.5, 0.3, h, size=12, color=PENT, font=BODYB, bold=True)
    text(s, x + 0.5, y, 5.3, 0.85, [[(a + " ", {"font": BODYB, "bold": True, "size": 10.5})], [(b, {"color": MUTED, "size": 9.5})]], size=10.5, line=1.0)
source(s, "Magnetite benchmark: Korshunov et al. 2025 (CC BY 4.0) report magnetite IoU 0.650 on LumenStone, the hardest of their ten classes. Phantom result: experiments/001-week1-gate.")
pic(s, G["lens_a"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_14")

# ---- 15 COMPLIANCE
s = new()
brand(s); pagenum(s, 15)
eyebrow(s, 0.6, 0.85, "Compliance, licences and safety")
headline(s, "Built to pass a mine's procurement,\nsafety and audit reviews.", y=1.15, size=32)
comp = [("POPIA", "Personal information is limited to accounts and notes; tenant isolation; nothing goes to an AI provider without explicit consent.", PYRR),
        ("Mine Health & Safety Act · OT safety", "Advisory first. A setpoint changes only with site authorisation, interlocks and rollback; OPC UA hardened to IEC 62443.", CHALC),
        ("SAMREC Code", "Outputs are 2D image fractions for process decisions — never public Mineral Resource or Reserve estimates.", PENT),
        ("MPRDA & data rights", "The mine owns its samples and data; models trained on a site's data stay licensed to that site.", MAG),
        ("Open-source licences", "Permissive-only shipped stack with an SBOM; GPL code excluded; dataset terms cited, assignable to Mintek.", INK),
        ("IP & originality", "Two clean git histories, dated failures kept, Mintek invention-credit process respected, AI assistance disclosed.", COPPER)]
for i, (a, b, c) in enumerate(comp):
    x = 0.6 + (i % 3) * 4.08
    y = 2.55 + (i // 3) * 2.15
    box(s, x, y, 3.85, 1.95, fill=PAPER, radius=0.12)
    box(s, x + 0.22, y + 0.22, 0.3, 0.3, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.22, y + 0.62, 3.45, 0.35, a, size=13.5, font=BODYB, bold=True)
    text(s, x + 0.22, y + 0.98, 3.45, 0.95, b, size=10.5, color=MUTED, line=1.05)
pic(s, G["lens_orig"], 11.95, 0.7, w=0.8, name="!!lens")
morph(s)
notes(s, "SCRIPT_15")

# ---- 16 THE ASK
s = new(dark=True)
pic(s, G["bg_close"], 0, 0, w=SW, h=SH)
brand(s, dark=True); pagenum(s, 16, dark=True)
pic(s, G["lens_title"], 8.6, 1.4, w=3.9, name="!!lens")
eyebrow(s, 0.7, 1.4, "Our ask", color=RGBColor(0xF2, 0x9A, 0x6A))
text(s, 0.7, 1.8, 7.7, 2.2, "One concentrator. One mineralogist. QEMSCAN-labelled sections. Twelve weeks.",
     size=34, color=WHITE, font=HEAD, bold=True, line=1.02, name="!!headline")
text(s, 0.7, 4.25, 7.4, 1.5, ["In return, at week 12, you get measured answers:",
                              "turnaround per decision · accuracy on your ore · refusal and coverage rates · a value estimate on your numbers."],
     size=15, color=RGBColor(0xDD, 0xE5, 0xEE), spacing=6, line=1.1)
text(s, 0.7, 5.95, 7.6, 0.5, [[("REEFPRINT", {"font": BODYB, "bold": True, "color": WHITE}), ("  otherwise known as KHANYA  ·  Team Sonar", {"color": RGBColor(0xC9, 0xD5, 0xE3)})]], size=15)
tag(s, 0.7, 6.85, "ILLUSTRATIVE")
text(s, 2.1, 6.86, 6, 0.3, "Background: Mixkit stock footage.", size=8.5, color=RGBColor(0x9F, 0xAE, 0xBF))
morph(s)
notes(s, "SCRIPT_16")

# ---- APPENDIX A: SOURCES
s = new()
brand(s); pagenum(s, 17)
eyebrow(s, 0.6, 0.85, "Appendix A · Sources and evidence")
headline(s, "Every number, and where it came from.", y=1.15, size=28)
srcs = [
    "Mintek–SCi Grad Hackathon 2026 — Problem 3 brief (“can take days”; deliverables) and organiser email (presentation structure).",
    "USGS, Mineral Commodity Summaries 2025, Platinum-group metals: SA platinum mine production 120,000 kg of world 170,000 kg (2024e); SA PGM reserves 63,000,000 kg of world >81,000,000 kg.",
    "Department of Mineral and Petroleum Resources, Mining Sector Performance 2024 (via Mining Review): mining 6.1% of nominal GDP; 474,736 direct employees.",
    "Valterra Platinum, H1 2026 short-form results (29 July 2026): rand basket price R45,993 per PGM ounce sold.",
    "Saskatchewan Research Council, Advanced Microanalysis Centre, QEMSCAN® price list (April 2017): modal mineralogy with liberation, $1,500 per sample.",
    "ALS Global, Mineralogy FAQ: turnaround is workload-dependent; a 1-week expedited option is available.",
    "Korshunov et al. 2025 (CC BY 4.0), LumenStone segmentation: magnetite IoU 0.650, the lowest of ten classes; polarised (XPL) channels improve segmentation.",
    "REEFPRINT/KHANYA repository: reports/ACCURACY-REPORT.md; reports/native-selected-test-20261001/ (candidate 42646cfa); live app capture 1 Oct 2026, result 98.2 s runtime.",
    "Price, pilot budget and value examples: team planning hypotheses (handover/PILOT-AND-BUSINESS.md), tagged ASSUMPTION or ILLUSTRATIVE wherever shown."]
text(s, 0.6, 2.2, 12.1, 4.8, srcs, size=11.5, color=INK, spacing=7, line=1.05)
notes(s, "Appendix — use in Q&A only.")
morph(s)

# ---- APPENDIX B: AI USE & ORIGINALITY
s = new()
brand(s); pagenum(s, 18)
eyebrow(s, 0.6, 0.85, "Appendix B · Originality and AI-assistance disclosure")
headline(s, "What is ours, what was assisted, and what is stock.", y=1.15, size=28)
items = [("Ours", "Problem framing, the refusal-by-design approach, model training and evaluation, the app, all measurements and every failure log — dated in two git histories."),
         ("AI-assisted", "Coding assistants (Claude, Codex) helped write code, documents, this deck's layout and the demo video's edit. Every number was checked against its source."),
         ("Synthetic", "The demo narration is a synthetic voice (en-GB); the music is an original composition generated in code."),
         ("Stock", "Mine, lab, bakkie, office and plant scenes are Mixkit stock footage under its free licence, labelled illustrative on screen."),
         ("Not claimed", "No live plant connection, no measured recovery gain, no UG2 or chromite accuracy, no patent or freedom-to-operate clearance.")]
for i, (a, b) in enumerate(items):
    y = 2.25 + i * 0.92
    box(s, 0.6, y, 12.1, 0.8, fill=PAPER, radius=0.1)
    text(s, 0.85, y + 0.22, 2.0, 0.4, a, size=14, font=BODYB, bold=True, color=[PYRR, PENT, CHALC, MAG, COPPER][i])
    text(s, 2.85, y + 0.13, 9.6, 0.62, b, size=11.5, color=INK, line=1.05)
notes(s, "Appendix — use in Q&A only.")
morph(s)

# ---------------------------------------------------------------- notes from the script file
script_path = os.path.join(D, "script_notes.json")
if os.path.exists(script_path):
    SC = json.load(open(script_path, encoding="utf-8"))
    for i, sl in enumerate(prs.slides, 1):
        tf = sl.notes_slide.notes_text_frame
        key = f"SCRIPT_{i}"
        if tf.text.strip() == key and key in SC:
            tf.text = SC[key]

name = os.path.join(OUT, "REEFPRINT-KHANYA-Team-Sonar-pitch.pptx")
prs.save(name)
print("saved", name, len(prs.slides), "slides")
